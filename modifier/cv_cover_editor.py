"""
AI-Powered Document Modifier
Generates tailored resumes and cover letters using AI
"""

import os
import re
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple, Optional
import openai
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn

from core.config import UserProfile, AIConfig
from core.database import JobListing

logger = logging.getLogger(__name__)

class AIDocumentModifier:
    """AI-powered document generation and modification"""
    
    def __init__(self, user_profile: UserProfile, ai_config: AIConfig):
        self.user_profile = user_profile
        self.ai_config = ai_config
        self.templates_dir = Path("templates")
        self.documents_dir = Path("documents")
        
        # Create directories
        self.templates_dir.mkdir(exist_ok=True)
        self.documents_dir.mkdir(exist_ok=True)
        
        # Initialize OpenAI
        if ai_config.openai_api_key:
            openai.api_key = ai_config.openai_api_key
        
        # Load or create templates
        self._ensure_templates_exist()
    
    def _ensure_templates_exist(self):
        """Create default templates if they don't exist"""
        resume_template_path = self.templates_dir / "resume_template.txt"
        cover_letter_template_path = self.templates_dir / "cover_letter_template.txt"
        
        if not resume_template_path.exists():
            self._create_default_resume_template(resume_template_path)
        
        if not cover_letter_template_path.exists():
            self._create_default_cover_letter_template(cover_letter_template_path)
    
    def _create_default_resume_template(self, path: Path):
        """Create a default resume template"""
        template = f"""
{self.user_profile.name.upper()}
{self.user_profile.email} | {self.user_profile.phone} | {self.user_profile.location}

PROFESSIONAL SUMMARY
Experienced {', '.join(self.user_profile.target_positions)} with {self.user_profile.experience_years} years of expertise in software development. Proficient in {', '.join(self.user_profile.skills[:5])} with a strong background in full-stack development and modern technologies.

TECHNICAL SKILLS
• Programming Languages: {', '.join([s for s in self.user_profile.skills if s.lower() in ['java', 'python', 'javascript', 'typescript', 'c++', 'c#']])}
• Frameworks & Libraries: {', '.join([s for s in self.user_profile.skills if s.lower() in ['spring', 'react', 'angular', 'vue', 'django', 'flask']])}
• Databases: {', '.join([s for s in self.user_profile.skills if s.lower() in ['mysql', 'postgresql', 'mongodb', 'redis']])}
• Tools & Technologies: {', '.join([s for s in self.user_profile.skills if s.lower() in ['docker', 'kubernetes', 'git', 'jenkins', 'aws']])}

PROFESSIONAL EXPERIENCE

Senior Software Developer | Current Company | 2021 - Present
• Developed and maintained enterprise-level applications serving 10,000+ users
• Led team of 3 developers in implementing microservices architecture
• Improved application performance by 40% through code optimization
• Collaborated with cross-functional teams to deliver projects on time

Software Developer | Previous Company | 2019 - 2021
• Built responsive web applications using modern JavaScript frameworks
• Integrated third-party APIs and payment gateways
• Participated in code reviews and maintained high code quality standards
• Mentored junior developers and conducted technical interviews

EDUCATION
{self.user_profile.education} in Computer Science
University Name | Year

PROJECTS
• E-commerce Platform: Full-stack web application with payment integration
• Task Management System: React-based SPA with real-time updates
• Data Analytics Dashboard: Python-based dashboard with data visualization

CERTIFICATIONS
• AWS Certified Developer
• Oracle Java Certification
• Scrum Master Certification
"""
        
        with open(path, 'w', encoding='utf-8') as f:
            f.write(template.strip())
        
        logger.info(f"📄 Created default resume template at {path}")
    
    def _create_default_cover_letter_template(self, path: Path):
        """Create a default cover letter template"""
        template = f"""
{self.user_profile.name}
{self.user_profile.email}
{self.user_profile.phone}
{self.user_profile.location}

[Date]

[Hiring Manager Name]
[Company Name]
[Company Address]

Dear Hiring Manager,

I am writing to express my strong interest in the [Job Title] position at [Company Name]. With {self.user_profile.experience_years} years of experience in software development and expertise in {', '.join(self.user_profile.skills[:3])}, I am excited about the opportunity to contribute to your team.

In my current role, I have successfully:
• Developed scalable applications using modern technologies
• Collaborated with cross-functional teams to deliver high-quality solutions
• Implemented best practices in software development and testing

I am particularly drawn to [Company Name] because of [Company Specific Reason]. Your commitment to [Company Value] aligns perfectly with my professional values and career goals.

I would welcome the opportunity to discuss how my skills in {', '.join(self.user_profile.skills[:2])} and passion for technology can contribute to [Company Name]'s continued success.

Thank you for your time and consideration. I look forward to hearing from you.

Sincerely,
{self.user_profile.name}
"""
        
        with open(path, 'w', encoding='utf-8') as f:
            f.write(template.strip())
        
        logger.info(f"📄 Created default cover letter template at {path}")
    
    def analyze_job_requirements(self, job: JobListing) -> Dict:
        """Use AI to analyze job requirements and extract key information"""
        prompt = f"""
        Analyze this job posting and extract structured information:
        
        Job Title: {job.title}
        Company: {job.company}
        Location: {job.location}
        Description: {job.description[:1500]}
        Requirements: {job.requirements[:1000]}
        
        Extract and return JSON with:
        {{
            "must_have_skills": ["skill1", "skill2"],
            "nice_to_have_skills": ["skill1", "skill2"],
            "experience_level": "junior/mid/senior",
            "required_years": number,
            "key_responsibilities": ["resp1", "resp2"],
            "company_culture": "description",
            "technical_stack": ["tech1", "tech2"],
            "soft_skills": ["skill1", "skill2"],
            "industry_keywords": ["keyword1", "keyword2"]
        }}
        
        Focus on technical requirements for software development roles.
        """
        
        try:
            if not self.ai_config.openai_api_key:
                logger.warning("No OpenAI API key - using fallback analysis")
                return self._fallback_job_analysis(job)
            
            response = openai.ChatCompletion.create(
                model=self.ai_config.model_name,
                messages=[
                    {"role": "system", "content": "You are an expert HR analyst specializing in tech job requirements. Return only valid JSON."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=1000,
                temperature=0.3
            )
            
            analysis = json.loads(response.choices[0].message.content.strip())
            logger.debug(f"✅ AI job analysis completed for {job.company}")
            return analysis
            
        except json.JSONDecodeError:
            logger.warning("Invalid JSON from AI - using fallback")
            return self._fallback_job_analysis(job)
        except Exception as e:
            logger.error(f"AI analysis failed: {e}")
            return self._fallback_job_analysis(job)
    
    def _fallback_job_analysis(self, job: JobListing) -> Dict:
        """Fallback job analysis using pattern matching"""
        full_text = f"{job.title} {job.description} {job.requirements}".lower()
        
        # Common tech skills
        tech_skills = []
        skill_patterns = {
            'java': r'\bjava\b',
            'python': r'\bpython\b',
            'javascript': r'\bjavascript\b|\bjs\b',
            'react': r'\breact\b',
            'angular': r'\bangular\b',
            'vue': r'\bvue\b',
            'spring': r'\bspring\b',
            'hibernate': r'\bhibernate\b',
            'mysql': r'\bmysql\b',
            'postgresql': r'\bpostgresql\b|\bpostgres\b',
            'docker': r'\bdocker\b',
            'kubernetes': r'\bkubernetes\b|\bk8s\b',
            'aws': r'\baws\b|\bamazon web services\b',
            'git': r'\bgit\b'
        }
        
        for skill, pattern in skill_patterns.items():
            if re.search(pattern, full_text):
                tech_skills.append(skill)
        
        # Experience level
        experience_level = "mid"
        if any(term in full_text for term in ['senior', 'lead', 'principal', '5+ years', '6+ years']):
            experience_level = "senior"
        elif any(term in full_text for term in ['junior', 'entry', '0-2 years', 'graduate']):
            experience_level = "junior"
        
        # Required years
        required_years = 3
        exp_match = re.search(r'(\d+)\+?\s*years?', full_text)
        if exp_match:
            required_years = int(exp_match.group(1))
        
        return {
            "must_have_skills": tech_skills[:5],
            "nice_to_have_skills": tech_skills[5:8],
            "experience_level": experience_level,
            "required_years": required_years,
            "key_responsibilities": ["Software development", "Code review", "Testing"],
            "company_culture": "Technology-focused",
            "technical_stack": tech_skills,
            "soft_skills": ["Communication", "Problem-solving", "Teamwork"],
            "industry_keywords": ["development", "software", "programming"]
        }
    
    def generate_tailored_resume(self, job: JobListing, job_analysis: Dict) -> str:
        """Generate a tailored resume for the specific job"""
        # Read base resume template
        try:
            with open(self.user_profile.base_resume_path, 'r', encoding='utf-8') as f:
                base_resume = f.read()
        except FileNotFoundError:
            logger.warning("Base resume template not found, using default")
            base_resume = self._get_default_resume_content()
        
        # Create AI prompt for resume tailoring
        prompt = f"""
        Tailor this resume for the specific job application:
        
        Base Resume:
        {base_resume}
        
        Job Details:
        - Title: {job.title}
        - Company: {job.company}
        - Required Skills: {', '.join(job_analysis.get('must_have_skills', []))}
        - Experience Level: {job_analysis.get('experience_level', 'mid')}
        - Key Responsibilities: {', '.join(job_analysis.get('key_responsibilities', []))}
        
        User Profile:
        - Skills: {', '.join(self.user_profile.skills)}
        - Experience: {self.user_profile.experience_years} years
        - Target Roles: {', '.join(self.user_profile.target_positions)}
        
        Instructions:
        1. Rewrite the professional summary to target this specific role
        2. Reorder and emphasize skills that match job requirements
        3. Adjust experience descriptions to highlight relevant achievements
        4. Use keywords from the job posting naturally
        5. Maintain professional formatting and ATS compatibility
        6. Keep the same structure but optimize content
        
        Return only the complete tailored resume.
        """
        
        try:
            if not self.ai_config.openai_api_key:
                logger.warning("No OpenAI API key - using template modification")
                return self._modify_resume_template(base_resume, job, job_analysis)
            
            response = openai.ChatCompletion.create(
                model=self.ai_config.model_name,
                messages=[
                    {"role": "system", "content": "You are a professional resume writer specializing in tech roles. Create ATS-optimized resumes."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=2000,
                temperature=0.4
            )
            
            tailored_resume = response.choices[0].message.content.strip()
            
            # Save to file
            resume_filename = f"resume_{job.company.replace(' ', '_')}_{job.id[:8]}.txt"
            resume_path = self.documents_dir / resume_filename
            
            with open(resume_path, 'w', encoding='utf-8') as f:
                f.write(tailored_resume)
            
            logger.info(f"✅ Tailored resume generated: {resume_filename}")
            return str(resume_path)
            
        except Exception as e:
            logger.error(f"AI resume generation failed: {e}")
            return self._modify_resume_template(base_resume, job, job_analysis)
    
    def _modify_resume_template(self, base_resume: str, job: JobListing, job_analysis: Dict) -> str:
        """Fallback resume modification without AI"""
        modified_resume = base_resume
        
        # Replace placeholders
        replacements = {
            '[Job Title]': job.title,
            '[Company Name]': job.company,
            '[Required Skills]': ', '.join(job_analysis.get('must_have_skills', [])[:3])
        }
        
        for placeholder, replacement in replacements.items():
            modified_resume = modified_resume.replace(placeholder, replacement)
        
        # Emphasize matching skills in summary
        required_skills = job_analysis.get('must_have_skills', [])
        user_matching_skills = [skill for skill in self.user_profile.skills 
                               if any(req.lower() in skill.lower() for req in required_skills)]
        
        if user_matching_skills:
            # Update professional summary
            summary_section = f"Experienced {job.title} with {self.user_profile.experience_years} years of expertise in {', '.join(user_matching_skills[:3])}."
            modified_resume = re.sub(
                r'PROFESSIONAL SUMMARY\n.*?\n\n',
                f'PROFESSIONAL SUMMARY\n{summary_section}\n\n',
                modified_resume,
                flags=re.DOTALL
            )
        
        # Save modified resume
        resume_filename = f"resume_{job.company.replace(' ', '_')}_{job.id[:8]}_template.txt"
        resume_path = self.documents_dir / resume_filename
        
        with open(resume_path, 'w', encoding='utf-8') as f:
            f.write(modified_resume)
        
        logger.info(f"📄 Template-based resume generated: {resume_filename}")
        return str(resume_path)
    
    def generate_cover_letter(self, job: JobListing, job_analysis: Dict) -> str:
        """Generate a tailored cover letter"""
        # Research company information
        company_info = self._research_company(job.company)
        
        # Read cover letter template
        try:
            with open(self.user_profile.cover_letter_template_path, 'r', encoding='utf-8') as f:
                base_template = f.read()
        except FileNotFoundError:
            logger.warning("Cover letter template not found, using default")
            base_template = self._get_default_cover_letter_content()
        
        prompt = f"""
        Write a compelling cover letter for this job application:
        
        Job Details:
        - Title: {job.title}
        - Company: {job.company}
        - Location: {job.location}
        - Key Requirements: {', '.join(job_analysis.get('must_have_skills', []))}
        - Company Culture: {job_analysis.get('company_culture', 'Not specified')}
        
        Candidate Profile:
        - Name: {self.user_profile.name}
        - Experience: {self.user_profile.experience_years} years
        - Skills: {', '.join(self.user_profile.skills[:5])}
        - Location: {self.user_profile.location}
        
        Company Research:
        {company_info}
        
        Template Structure:
        {base_template}
        
        Instructions:
        1. Write a personalized opening that shows genuine interest
        2. Highlight 2-3 most relevant achievements/skills
        3. Show knowledge of the company and role
        4. Explain why you're interested in this specific position
        5. Professional but engaging tone
        6. Keep to 3-4 paragraphs
        7. Include specific examples of relevant experience
        
        Return only the complete cover letter.
        """
        
        try:
            if not self.ai_config.openai_api_key:
                logger.warning("No OpenAI API key - using template modification")
                return self._modify_cover_letter_template(base_template, job, job_analysis)
            
            response = openai.ChatCompletion.create(
                model=self.ai_config.model_name,
                messages=[
                    {"role": "system", "content": "You are a professional career coach specializing in tech industry cover letters. Write compelling, personalized cover letters."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=1500,
                temperature=0.6
            )
            
            cover_letter = response.choices[0].message.content.strip()
            
            # Save to file
            letter_filename = f"cover_letter_{job.company.replace(' ', '_')}_{job.id[:8]}.txt"
            letter_path = self.documents_dir / letter_filename
            
            with open(letter_path, 'w', encoding='utf-8') as f:
                f.write(cover_letter)
            
            logger.info(f"✅ Cover letter generated: {letter_filename}")
            return str(letter_path)
            
        except Exception as e:
            logger.error(f"AI cover letter generation failed: {e}")
            return self._modify_cover_letter_template(base_template, job, job_analysis)
    
    def _modify_cover_letter_template(self, template: str, job: JobListing, job_analysis: Dict) -> str:
        """Fallback cover letter modification without AI"""
        current_date = datetime.now().strftime("%B %d, %Y")
        
        # Replace template placeholders
        replacements = {
            '[Date]': current_date,
            '[Job Title]': job.title,
            '[Company Name]': job.company,
            '[Company Specific Reason]': f"your innovative work in {job_analysis.get('industry_keywords', ['technology'])[0]}",
            '[Company Value]': "innovation and excellence",
            '[Hiring Manager Name]': "Hiring Manager"
        }
        
        modified_letter = template
        for placeholder, replacement in replacements.items():
            modified_letter = modified_letter.replace(placeholder, replacement)
        
        # Save modified cover letter
        letter_filename = f"cover_letter_{job.company.replace(' ', '_')}_{job.id[:8]}_template.txt"
        letter_path = self.documents_dir / letter_filename
        
        with open(letter_path, 'w', encoding='utf-8') as f:
            f.write(modified_letter)
        
        logger.info(f"📄 Template-based cover letter generated: {letter_filename}")
        return str(letter_path)
    
    def _research_company(self, company_name: str) -> str:
        """Basic company research (placeholder for future enhancement)"""
        # This could be enhanced with web scraping or API calls
        return f"Research shows that {company_name} is a leading company in the technology sector, known for innovation and employee development."
    
    def generate_documents_for_job(self, job: JobListing) -> Tuple[Optional[str], Optional[str]]:
        """Generate both resume and cover letter for a job"""
        logger.info(f"📝 Generating documents for {job.title} at {job.company}")
        
        try:
            # Analyze job requirements
            job_analysis = self.analyze_job_requirements(job)
            
            # Generate tailored resume
            resume_path = self.generate_tailored_resume(job, job_analysis)
            
            # Generate cover letter
            cover_letter_path = self.generate_cover_letter(job, job_analysis)
            
            logger.info(f"✅ Documents generated successfully for {job.company}")
            return resume_path, cover_letter_path
            
        except Exception as e:
            logger.error(f"Document generation failed for {job.company}: {e}")
            return None, None
    
    def create_docx_resume(self, text_resume_path: str) -> str:
        """Convert text resume to professional DOCX format"""
        try:
            # Read text resume
            with open(text_resume_path, 'r', encoding='utf-8') as f:
                resume_text = f.read()
            
            # Create Word document
            doc = Document()
            
            # Set document margins
            sections = doc.sections
            for section in sections:
                section.top_margin = Inches(0.5)
                section.bottom_margin = Inches(0.5)
                section.left_margin = Inches(0.75)
                section.right_margin = Inches(0.75)
            
            # Parse and format resume sections
            lines = resume_text.split('\n')
            current_section = None
            
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                # Check if it's a section header (all caps or has specific keywords)
                if (line.isupper() and len(line) > 3) or line.startswith('PROFESSIONAL') or line.startswith('TECHNICAL') or line.startswith('EDUCATION'):
                    # Section header
                    p = doc.add_paragraph()
                    run = p.add_run(line)
                    run.font.size = Pt(12)
                    run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    current_section = line
                elif line.startswith('•') or line.startswith('-'):
                    # Bullet point
                    p = doc.add_paragraph(line, style='List Bullet')
                    p.paragraph_format.left_indent = Inches(0.25)
                elif '|' in line and current_section != 'CONTACT':
                    # Job title with company and dates
                    p = doc.add_paragraph()
                    run = p.add_run(line)
                    run.font.size = Pt(11)
                    run.bold = True
                else:
                    # Regular paragraph
                    p = doc.add_paragraph(line)
                    if current_section is None:  # Name at top
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.runs[0]
                        run.font.size = Pt(16)
                        run.bold = True
            
            # Save DOCX file
            docx_path = text_resume_path.replace('.txt', '.docx')
            doc.save(docx_path)
            
            logger.info(f"📄 DOCX resume created: {docx_path}")
            return docx_path
            
        except Exception as e:
            logger.error(f"DOCX creation failed: {e}")
            return text_resume_path
    
    def _get_default_resume_content(self) -> str:
        """Get default resume content if template is missing"""
        return f"""
{self.user_profile.name.upper()}
{self.user_profile.email} | {self.user_profile.phone} | {self.user_profile.location}

PROFESSIONAL SUMMARY
Experienced software developer with {self.user_profile.experience_years} years of expertise in modern web technologies.

TECHNICAL SKILLS
• {', '.join(self.user_profile.skills)}

PROFESSIONAL EXPERIENCE
Software Developer | Company Name | Years
• Developed applications using modern technologies
• Collaborated with cross-functional teams

EDUCATION
{self.user_profile.education}
"""
    
    def _get_default_cover_letter_content(self) -> str:
        """Get default cover letter content if template is missing"""
        return f"""
{self.user_profile.name}
{self.user_profile.email}
{self.user_profile.phone}

[Date]

Dear Hiring Manager,

I am writing to express my interest in the [Job Title] position at [Company Name]. With {self.user_profile.experience_years} years of experience in software development, I am excited about this opportunity.

I believe my skills in {', '.join(self.user_profile.skills[:3])} make me a strong candidate for this role.

Thank you for your consideration.

Sincerely,
{self.user_profile.name}
"""


# Legacy function for backwards compatibility
def generate_documents(job):
    """Legacy function for backwards compatibility"""
    from core.config import load_config
    
    try:
        config = load_config()
        modifier = AIDocumentModifier(config.user_profile, config.ai)
        
        # Convert dict to JobListing if needed
        if isinstance(job, dict):
            from core.database import JobListing
            job_listing = JobListing(
                id=job.get('id', ''),
                title=job.get('title', ''),
                company=job.get('company', ''),
                location=job.get('location', ''),
                salary=job.get('salary', ''),
                description=job.get('description', ''),
                requirements=job.get('requirements', ''),
                url=job.get('url', ''),
                source=job.get('source', ''),
                date_scraped=job.get('date_scraped', ''),
                country=job.get('country', ''),
                remote=job.get('remote', False),
                visa_sponsorship=job.get('visa_sponsorship', False)
            )
        else:
            job_listing = job
        
        resume_path, cover_letter_path = modifier.generate_documents_for_job(job_listing)
        
        if resume_path and cover_letter_path:
            print(f"✅ [Modifier] Documents generated for: {job_listing.title}")
        else:
            print(f"❌ [Modifier] Failed to generate documents for: {job_listing.title}")
            
    except Exception as e:
        print(f"❌ [Modifier] Error: {e}")
        logger.error(f"Legacy document generation failed: {e}")